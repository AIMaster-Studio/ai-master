"""Prepare review copies, never overwrite original PPTX. Requires python-pptx."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
ROOT=Path(__file__).resolve().parent.parent
OUT=ROOT/'.local/submission-drafts'
REPLACEMENTS={
 '公开仓库：AIMaster-Studio/ai-master':'项目仓库地址：盲审草稿不展示',
 'AIMaster-Studio/ai-master':'项目仓库（盲审草稿隐去）',
 'AIMaster-Studio':'项目团队',
 '21 组双盲测试':'21 条历史探索性样本（标注流程待复核）',
 '零假阳性（FP = 0）——不会错误通过无效讲解':'样本内 FP = 0；不代表系统不存在错误放行',
 '线上实际使用：deepseek-flash':'历史演示配置：deepseek-flash；当前配置待核验',
 '准确率 90.4762%':'历史小样本准确率 90.4762%（非稳定性能）'
}
def visit(shapes):
 for shape in shapes:
  if hasattr(shape,'shapes'):yield from visit(shape.shapes)
  if shape.has_text_frame:yield shape.text_frame
  if shape.has_table:
   for row in shape.table.rows:
    for cell in row.cells:yield cell.text_frame

def clean(frame):
 for paragraph in frame.paragraphs:
  old=paragraph.text;new=old
  for a,b in REPLACEMENTS.items():new=new.replace(a,b)
  if old!=new and paragraph.runs:
   paragraph.runs[0].text=new
   for run in paragraph.runs[1:]:run.text=''
  for run in paragraph.runs:
   url=run.hyperlink.address
   if url and any(x in url.lower() for x in ['github.com','aimaster-studio']):run.hyperlink.address=None

def main():
 OUT.mkdir(parents=True,exist_ok=True)
 for source in sorted(ROOT.glob('*.pptx')):
  prs=Presentation(source)
  for slide in prs.slides:
   original=' '.join(f.text for f in visit(slide.shapes))
   for frame in visit(slide.shapes):clean(frame)
   if slide.has_notes_slide:clean(slide.notes_slide.notes_text_frame)
   foot=slide.shapes.add_textbox(Inches(.25),prs.slide_height-Inches(.26),prs.slide_width-Inches(.5),Inches(.22))
   p=foot.text_frame.paragraphs[0];p.text='AI辅助制作 · 审核草稿（图片/二维码/匿名要求待人工复核）';p.font.size=Pt(9);p.font.color.rgb=RGBColor(110,110,110)
   if '90.4762' in original:
    box=slide.shapes.add_textbox(Inches(.3),prs.slide_height-Inches(.85),prs.slide_width-Inches(.6),Inches(.5))
    box.fill.solid();box.fill.fore_color.rgb=RGBColor(255,248,225)
    p=box.text_frame.paragraphs[0];p.text='数据边界：21条历史探索性样本；测试模型不等同当前生产配置。样本内FP=0不代表零误判。';p.font.size=Pt(12);p.font.color.rgb=RGBColor(100,65,0)
  props=prs.core_properties;props.author='';props.last_modified_by='';props.comments='AI-assisted review draft; not certified for submission';props.keywords=''
  dest=OUT/(source.stem+'_审核草稿.pptx');prs.save(dest);print(dest.relative_to(ROOT).as_posix())
if __name__=='__main__':main()
